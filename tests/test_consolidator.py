import tempfile
import unittest
import shutil
import uuid
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from modules.afip_consolidator.controller import ConsolidatorController
from modules.afip_consolidator.processor import consolidar
from modules.afip_consolidator.exporter import exportar
from modules.afip_consolidator.models import ErrorConsolidacion, ModoSalida, COLUMNA_PERIODO


class ConsolidacionTests(unittest.TestCase):
    def setUp(self):
        raiz = Path(tempfile.gettempdir()).resolve()
        self.dir = raiz / ("consolidador-test-" + uuid.uuid4().hex)
        self.dir.mkdir()
        def limpiar():
            if self.dir.resolve().parent != raiz or not self.dir.name.startswith("consolidador-test-"):
                raise RuntimeError("Directorio de prueba fuera del área temporal.")
            shutil.rmtree(self.dir)
        self.addCleanup(limpiar)

    def csv(self, texto, nombre="entrada.csv", encoding="utf-8-sig"):
        ruta = self.dir / nombre
        ruta.write_text(texto, encoding=encoding, newline="")
        return ruta

    def excel(self, filas, nombre="entrada.xlsx"):
        ruta = self.dir / nombre
        libro = Workbook()
        for fila in filas:
            libro.active.append(fila)
        libro.save(ruta)
        libro.close()
        return ruta

    def test_conserva_csv_y_dos_salidas(self):
        entrada = self.csv('Reporte ficticio\nFecha;Tipo;CUIT;Importe;Detalle\n01/01/2026;001;00123456789;1.250,50;"Texto; con separador"\n01/02/2026;001;00123456789;error;=1+1\nfecha mala;001;;1,234.56;NA\n')
        resultado = consolidar([entrada])
        self.assertEqual(len(resultado.comprobantes), 3)
        self.assertEqual(resultado.comprobantes["Importe"].tolist(), ["1.250,50", "error", "1,234.56"])
        self.assertEqual(resultado.comprobantes["CUIT"].iloc[0], "00123456789")
        self.assertEqual(resultado.comprobantes[COLUMNA_PERIODO].iloc[2], "Sin fecha válida")
        self.assertTrue(resultado.avisos)
        for modo in ModoSalida:
            salida = self.dir / (modo.name + ".xlsx")
            exportar(resultado, salida, modo)
            libro = load_workbook(salida)
            try:
                hojas = libro.worksheets
                self.assertEqual(sum(h.max_row - 1 for h in hojas), 3)
                valores = [fila for h in hojas for fila in list(h.values)[1:]]
                self.assertEqual([f[3] for f in valores], ["1.250,50", "error", "1,234.56"])
                celda = hojas[0]["E3"] if modo == ModoSalida.SABANA else hojas[1]["E2"]
                self.assertEqual(celda.value, "=1+1")
                self.assertEqual(celda.data_type, "s")
            finally:
                libro.close()

    def test_numericos_excel_no_se_modifican(self):
        entrada = self.excel([["Reporte"], ["Fecha", "Tipo", "CUIT", "Número", "Importe"],
                              ["01/01/2026", 1, 20123456789, 123, 1250.5],
                              ["02/01/2026", 1, None, None, None]])
        resultado = consolidar([entrada])
        self.assertEqual(resultado.comprobantes["CUIT"].iloc[0], 20123456789)
        self.assertEqual(resultado.comprobantes["Número"].iloc[0], 123)
        self.assertEqual(resultado.comprobantes["Importe"].iloc[0], 1250.5)
        exportar(resultado, self.dir / "salida.xlsx", ModoSalida.SABANA)
        libro = load_workbook(self.dir / "salida.xlsx")
        self.assertEqual(libro.active["C2"].value, 20123456789)
        libro.close()

    def test_separadores_y_cp1252(self):
        for sep in [",", ";", "\t"]:
            for encoding in ["utf-8-sig", "cp1252"]:
                entrada = self.csv(sep.join(["Fecha", "Tipo", "Detalle"]) + "\n" +
                                   sep.join(["01/01/2026", "001", "Peña"]) + "\n", encoding=encoding)
                self.assertEqual(consolidar([entrada]).comprobantes["Detalle"].iloc[0], "Peña")

    def test_columnas_distintas_se_conservan(self):
        a = self.csv("Fecha;Tipo;A\n01/01/2026;001;x\n", "a.csv")
        b = self.csv("Fecha;Tipo;B\n02/01/2026;001;y\n", "b.csv")
        resultado = consolidar([a, b])
        self.assertEqual(len(resultado.comprobantes), 2)
        self.assertIn("B", resultado.comprobantes.columns)
        self.assertTrue(resultado.avisos)

    def test_filas_duplicadas_se_conservan_seleccion_repetida_no(self):
        entrada = self.csv("Fecha;Tipo\n01/01/2026;001\n01/01/2026;001\n")
        self.assertEqual(len(consolidar([entrada, entrada]).comprobantes), 2)

    def test_error_invalida_resultado_previo(self):
        entrada = self.csv("Fecha;Tipo\n01/01/2026;001\n")
        controller = ConsolidatorController()
        controller.seleccionar([entrada])
        controller.preparar()
        entrada.write_text("archivo roto", encoding="utf-8")
        with self.assertRaises(ErrorConsolidacion):
            controller.preparar()
        self.assertIsNone(controller.resultado)
        with self.assertRaises(ValueError):
            controller.exportar(self.dir / "salida.xlsx", ModoSalida.SABANA)

    def test_no_exporta_parcial_si_un_archivo_falla(self):
        bueno = self.csv("Fecha;Tipo\n01/01/2026;001\n", "bueno.csv")
        malo = self.csv("archivo roto", "malo.csv")
        with self.assertRaisesRegex(ErrorConsolidacion, "malo.csv"):
            consolidar([bueno, malo])

    def test_rechaza_cabecera_duplicada_o_sin_nombre(self):
        for texto in ["Fecha;Tipo;Tipo\n01/01/2026;1;1\n", "Fecha;Tipo;\n01/01/2026;1;dato\n"]:
            with self.assertRaises(ErrorConsolidacion):
                consolidar([self.csv(texto)])

    def test_filas_vacias_y_trazabilidad(self):
        entrada = self.csv("Titulo\nFecha;Tipo\n01/01/2026;001\n;\nfecha mala;001\n")
        resultado = consolidar([entrada])
        self.assertEqual(resultado.comprobantes["Origen - fila"].tolist(), [3, 5])
        self.assertEqual(resultado.archivos[0].filas_vacias, 1)

    def test_protege_entrada_y_destino_existente(self):
        entrada = self.excel([["Fecha", "Tipo"], ["01/01/2026", 1]])
        resultado = consolidar([entrada])
        original = entrada.read_bytes()
        with self.assertRaises(ValueError):
            exportar(resultado, entrada, ModoSalida.SABANA, reemplazar=True)
        self.assertEqual(original, entrada.read_bytes())
        destino = self.dir / "existente.xlsx"
        destino.write_bytes(b"original")
        with self.assertRaises(FileExistsError):
            exportar(resultado, destino, ModoSalida.SABANA)
        self.assertEqual(destino.read_bytes(), b"original")

    def test_guardado_fallido_no_destruye_destino(self):
        entrada = self.csv("Fecha;Tipo\n01/01/2026;001\n")
        destino = self.dir / "destino.xlsx"
        destino.write_bytes(b"original")
        with patch("modules.afip_consolidator.exporter._formatear", side_effect=RuntimeError("fallo simulado")):
            with self.assertRaises(RuntimeError):
                exportar(consolidar([entrada]), destino, ModoSalida.SABANA, reemplazar=True)
        self.assertEqual(destino.read_bytes(), b"original")
        self.assertEqual(list(self.dir.glob(".consolidado-*")), [])

    def test_excel_formula_se_rechaza(self):
        entrada = self.excel([["Fecha", "Tipo", "Importe"], ["01/01/2026", 1, "=1+1"]])
        with self.assertRaisesRegex(ErrorConsolidacion, "fórmula"):
            consolidar([entrada])

    def test_multiples_hojas_se_rechazan(self):
        entrada = self.excel([["Fecha", "Tipo"], ["01/01/2026", 1]])
        libro = load_workbook(entrada)
        libro.create_sheet("Otra").append(["Dato"])
        libro.save(entrada)
        libro.close()
        with self.assertRaisesRegex(ErrorConsolidacion, "una hoja"):
            consolidar([entrada])

    def test_no_reimportar_salida(self):
        entrada = self.csv("Fecha;Tipo\n01/01/2026;001\n")
        salida = self.dir / "salida.xlsx"
        exportar(consolidar([entrada]), salida, ModoSalida.SABANA)
        with self.assertRaises(ErrorConsolidacion):
            consolidar([salida])

    def test_rechaza_texto_que_excel_truncaria(self):
        entrada = self.csv("Fecha;Tipo;Detalle\n01/01/2026;001;" + "x" * 32768 + "\n")
        with self.assertRaisesRegex(ErrorConsolidacion, "demasiado largo"):
            consolidar([entrada])

    def test_csv_comillas_y_salto_de_linea(self):
        entrada = self.csv('Fecha,Tipo,Detalle\n01/01/2026,001,"Texto, con coma\ny segunda línea"\n')
        resultado = consolidar([entrada])
        self.assertEqual(len(resultado.comprobantes), 1)
        self.assertEqual(resultado.comprobantes['Detalle'].iloc[0], 'Texto, con coma\ny segunda línea')

    def test_origen_inequivoco_entre_carpetas(self):
        entrada = self.csv("Fecha;Tipo\n01/01/2026;001\n")
        carpeta = self.dir / "otra"
        carpeta.mkdir()
        otra = carpeta / entrada.name
        otra.write_bytes(entrada.read_bytes())
        with self.assertRaisesRegex(ErrorConsolidacion, "mismo nombre"):
            consolidar([entrada, otra])

    def test_csv_afip_fecha_emision(self):
        entrada = self.csv('"Fecha de Emisión";"Tipo de Comprobante";"Nro. Doc. Comprador";"Importe Total"\n'
                           '"02/04/2026";"001";"00000000001";"1.234,56"\n',
                           'comprobantes_periodo_202604_ventas.csv', encoding='cp1252')
        r = consolidar([entrada], periodo=('04/2026', '04/2026'))
        self.assertEqual(r.comprobantes[COLUMNA_PERIODO].tolist(), ['2026-04'])
        self.assertEqual(r.comprobantes['Importe Total'].iloc[0], '1.234,56')
        self.assertIn('Falta Compras 04/2026', r.avisos)
        self.assertNotIn('Falta Ventas 04/2026', r.avisos)

    def test_orden_mes_compras_ventas_y_hojas_sin_revisiones(self):
        entradas = [self.csv(f'Fecha;Tipo;ID\n{fecha};001;{nombre}\n',nombre+'.csv')
                    for nombre,fecha in [('ventas_marzo','01/03/2024'),('ventas_febrero','01/02/2024'),
                                         ('compras_marzo','20/03/2024'),('compras_febrero','20/02/2024')]]
        r = consolidar(entradas,periodo=('02/2024','03/2024'))
        orden = ['compras_febrero','ventas_febrero','compras_marzo','ventas_marzo']
        self.assertEqual(r.comprobantes['ID'].tolist(), orden)
        self.assertFalse(any(a.startswith('Falta ') for a in r.avisos))
        for modo in ModoSalida:
            salida = self.dir / (modo.name+'.xlsx')
            exportar(r,salida,modo)
            w=load_workbook(salida)
            try:
                self.assertEqual(w.sheetnames,['Vista Sábana'] if modo==ModoSalida.SABANA else ['2024-02','2024-03'])
                self.assertEqual([f[2] for h in w for f in list(h.values)[1:]],orden)
            finally:
                w.close()

    def test_periodo_invalido_y_no_filtra_datos(self):
        entrada=self.csv('Fecha;Tipo\n01/01/2024;001\n','compras.csv')
        for periodo in [('13/2024','01/2025'),('02/2024','01/2024')]:
            with self.assertRaises(ErrorConsolidacion):
                consolidar([entrada],periodo=periodo)
        r=consolidar([entrada],periodo=('02/2024','02/2024'))
        self.assertEqual(len(r.comprobantes),1)
        self.assertIn('Falta Compras 02/2024',r.avisos)
        self.assertIn('Falta Ventas 02/2024',r.avisos)


if __name__ == "__main__":
    unittest.main()
