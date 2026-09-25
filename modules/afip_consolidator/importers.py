from __future__ import annotations

import csv
import io
import unicodedata
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from .models import (
    ArchivoOrigen, ErrorConsolidacion, COLUMNAS_ORIGEN,
    COLUMNA_ARCHIVO, COLUMNA_HOJA, COLUMNA_FILA, COLUMNA_PERIODO,
)


EXTENSIONES = {".xlsx", ".xlsm", ".xls", ".csv"}
FECHAS = {"fecha", "fecha de emision"}


def _normalizar_texto(valor: object) -> str:
    texto = unicodedata.normalize("NFKD", str(valor).strip().lower())
    return " ".join("".join(c for c in texto if not unicodedata.combining(c)).split())


def _vacio(valor: object) -> bool:
    return valor is None or (isinstance(valor, str) and not valor.strip()) or bool(pd.isna(valor))


def _encabezado(fila: list) -> bool:
    nombres = {_normalizar_texto(v) for v in fila if not _vacio(v)}
    return bool(FECHAS & nombres) and bool({"tipo", "tipo de comprobante"} & nombres)


def _fecha(valor: object) -> str | None:
    """Calcula solo el período; conserva el valor original en el reporte."""
    if isinstance(valor, (datetime, date, pd.Timestamp)):
        return valor.strftime("%Y-%m")
    if not isinstance(valor, str):
        return None  # No interpretar números como nanosegundos o fechas Excel.
    for formato in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(valor.strip(), formato).strftime("%Y-%m")
        except ValueError:
            continue
    return None


def _leer_csv(ruta: Path) -> list[list[str]]:
    contenido = ruta.read_bytes()
    try:
        texto = contenido.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            texto = contenido.decode("cp1252")
        except UnicodeDecodeError as exc:
            raise ErrorConsolidacion("CSV: utilice UTF-8 o Windows-1252.") from exc
    # La cabecera puede venir después de un título de longitud distinta.
    # Probar separadores explícitos evita que el título confunda al detector.
    candidatos = []
    for separador in (";", ",", "\t"):
        try:
            filas = list(csv.reader(io.StringIO(texto, newline=""), delimiter=separador, strict=True))
        except csv.Error:
            continue
        if any(_encabezado(f) for f in filas[:30]):
            candidatos.append(filas)
    if len(candidatos) != 1:
        raise ErrorConsolidacion(
            "CSV: no se identificó una cabecera y separador únicos. "
            "Use coma, punto y coma o tabulación, con Fecha y Tipo."
        )
    return candidatos[0]


def _leer_excel(ruta: Path) -> tuple[list[list], str]:
    if ruta.suffix.lower() == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise ErrorConsolidacion("Para abrir .xls instale las dependencias de requirements.txt.") from exc
        libro = xlrd.open_workbook(str(ruta))
        try:
            hojas = [s for s in libro.sheets() if s.nrows and s.ncols]
            if len(hojas) != 1:
                raise ErrorConsolidacion("Se requiere exactamente una hoja con datos por archivo de entrada.")
            hoja = hojas[0]
            filas = []
            for indice in range(hoja.nrows):
                valores = []
                for celda in hoja.row(indice):
                    if celda.ctype == xlrd.XL_CELL_ERROR:
                        raise ErrorConsolidacion(f"Error de Excel en fila {indice + 1}.")
                    valor = celda.value
                    if celda.ctype == xlrd.XL_CELL_DATE:
                        valor = xlrd.xldate_as_datetime(valor, libro.datemode)
                    valores.append(valor)
                filas.append(valores)
            return filas, hoja.name
        finally:
            libro.release_resources()
    libro = load_workbook(ruta, read_only=True, data_only=False)
    try:
        hojas = []
        for hoja in libro.worksheets:
            filas = list(hoja.iter_rows())
            if any(any(c.value is not None for c in fila) for fila in filas):
                hojas.append((hoja.title, filas))
        if len(hojas) != 1:
            raise ErrorConsolidacion("Se requiere exactamente una hoja con datos por archivo de entrada.")
        nombre, celdas = hojas[0]
        for fila in celdas:
            for celda in fila:
                if celda.data_type in {"f", "e"}:
                    raise ErrorConsolidacion(
                        f"La celda {celda.coordinate} contiene una fórmula o error. "
                        "Exporte una copia con valores antes de consolidar."
                    )
        return [[c.value for c in fila] for fila in celdas], nombre
    finally:
        libro.close()


def leer_archivo(ruta: Path) -> tuple[pd.DataFrame, ArchivoOrigen, list[str]]:
    if not ruta.is_file() or ruta.suffix.lower() not in EXTENSIONES:
        raise ErrorConsolidacion("Seleccione un archivo existente .xlsx, .xlsm, .xls o .csv.")
    if ruta.suffix.lower() == ".csv":
        filas, hoja = _leer_csv(ruta), "CSV"
    else:
        filas, hoja = _leer_excel(ruta)
    cabecera = next((i for i, fila in enumerate(filas[:30]) if _encabezado(fila)), None)
    if cabecera is None:
        raise ErrorConsolidacion("No se encontró Fecha y Tipo / Tipo de Comprobante en las primeras 30 filas.")
    ancho = max(len(f) for f in filas[cabecera:])
    filas = [f + [None] * (ancho - len(f)) for f in filas]
    # Solo omitir columnas totalmente vacías, incluido el encabezado.
    indices = [i for i in range(ancho) if any(not _vacio(f[i]) for f in filas[cabecera:])]
    nombres = [str(filas[cabecera][i]).strip() if not _vacio(filas[cabecera][i]) else "" for i in indices]
    normalizados = [_normalizar_texto(n) for n in nombres]
    if any(len(n) > 32767 for n in nombres):
        raise ErrorConsolidacion("Un encabezado excede los 32.767 caracteres que admite Excel.")
    if "" in nombres or len(set(normalizados)) != len(nombres):
        raise ErrorConsolidacion("Hay columnas sin título o con títulos duplicados; corríjalas en una copia.")
    if set(normalizados) & {_normalizar_texto(n) for n in COLUMNAS_ORIGEN}:
        raise ErrorConsolidacion("El archivo ya contiene columnas de origen: podría ser una salida anterior.")
    columna_fecha = next(i for i, nombre in enumerate(normalizados) if nombre in FECHAS)
    datos, numeros, periodos, invalidas = [], [], [], []
    vacias = 0
    for numero, fila in enumerate(filas[cabecera + 1:], start=cabecera + 2):
        valores = [fila[i] for i in indices]
        if any(isinstance(v, str) and len(v) > 32767 for v in valores):
            raise ErrorConsolidacion(f"La fila {numero} contiene texto demasiado largo para Excel (máximo 32.767 caracteres).")
        if all(_vacio(v) for v in valores):
            vacias += 1
            continue
        periodo = _fecha(valores[columna_fecha])
        if periodo is None:
            invalidas.append(numero)
        datos.append(valores)
        numeros.append(numero)
        periodos.append(periodo or "Sin fecha válida")
    if not datos:
        raise ErrorConsolidacion("El archivo no contiene filas de datos.")
    # object impide que un entero se convierta en float por otra celda vacía.
    tabla = pd.DataFrame(datos, columns=nombres, dtype=object)
    tabla[COLUMNA_ARCHIVO] = ruta.name
    tabla[COLUMNA_HOJA] = hoja
    tabla[COLUMNA_FILA] = numeros
    tabla[COLUMNA_PERIODO] = periodos
    avisos = []
    if invalidas:
        muestra = ", ".join(map(str, invalidas[:20]))
        avisos.append(f"{ruta.name}: {len(invalidas)} fila(s) sin fecha válida ({muestra}). "
                      "Se conservan en 'Sin fecha válida'; consulte Origen - fila.")
    if vacias:
        avisos.append(f"{ruta.name}: se omitieron {vacias} fila(s) completamente vacía(s).")
    return tabla, ArchivoOrigen(ruta.resolve(), len(datos), vacias, hoja), avisos
