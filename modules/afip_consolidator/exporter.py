from __future__ import annotations

import os
import tempfile
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Resultado, ModoSalida, COLUMNA_PERIODO


def _formatear(hoja) -> None:
    for fila in hoja:
        for celda in fila:
            # pandas/openpyxl interpretan '=' como fórmula. Estos campos son datos.
            if celda.data_type == "f":
                celda.data_type = "s"
    for celda in hoja[1]:
        celda.fill = PatternFill("solid", fgColor="1F4E78")
        celda.font = Font(bold=True, color="FFFFFF")
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    for indice, columna in enumerate(hoja.columns, 1):
        ancho = max((len(str(c.value)) if c.value is not None else 0 for c in columna), default=10)
        hoja.column_dimensions[get_column_letter(indice)].width = min(max(ancho + 2, 12), 48)


def exportar(resultado: Resultado, ruta_salida: str | Path, modo: ModoSalida,
             *, reemplazar: bool = False) -> Path:
    modo = ModoSalida(modo)
    ruta = Path(ruta_salida).resolve()
    if ruta.suffix.lower() != ".xlsx":
        raise ValueError("La salida debe tener extensión .xlsx.")
    for origen in resultado.archivos:
        if ruta == origen.ruta or (ruta.exists() and origen.ruta.exists() and os.path.samefile(ruta, origen.ruta)):
            raise ValueError("El archivo de salida no puede reemplazar un archivo de entrada.")
    if ruta.exists() and not reemplazar:
        raise FileExistsError("El destino ya existe. Elija otro nombre o confirme su reemplazo.")
    if resultado.comprobantes.empty:
        raise ValueError("No hay filas para exportar.")
    if len(resultado.comprobantes.columns) > 16384:
        raise ValueError("La cantidad de columnas excede el límite de Excel.")
    grupos = [("Vista Sábana", resultado.comprobantes)] if modo == ModoSalida.SABANA else list(
        resultado.comprobantes.groupby(COLUMNA_PERIODO, sort=True)
    )
    if any(len(tabla) > 1048575 for _, tabla in grupos):
        raise ValueError("Una hoja excede el límite de filas de Excel. Divida la selección.")
    # No crear carpetas implícitamente: el usuario selecciona una carpeta existente.
    descriptor, temporal = tempfile.mkstemp(prefix=".consolidado-", suffix=".xlsx", dir=ruta.parent)
    os.close(descriptor)
    try:
        with pd.ExcelWriter(temporal, engine="openpyxl") as writer:
            for nombre, tabla in grupos:
                tabla.to_excel(writer, sheet_name=str(nombre), index=False)
            for hoja in writer.book.worksheets:
                _formatear(hoja)
        if ruta.exists() and not reemplazar:
            raise FileExistsError("El destino apareció durante la exportación. Elija otro nombre.")
        os.replace(temporal, ruta)
    finally:
        if os.path.exists(temporal):
            os.unlink(temporal)
    return ruta
